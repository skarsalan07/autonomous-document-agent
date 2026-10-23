"""Streamlit interface for the autonomous document agent."""

import json
import logging
import re
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document

from app.agent import run_agent
from app.config import Settings, get_settings
from app.exceptions import ConfigurationError, PlanningError, RetryableOperationError
from app.schemas import AgentRequest, AgentResponse

logger = logging.getLogger(__name__)
WORKSPACE = Path(__file__).resolve().parent
FILENAME_PATTERN = re.compile(r"^document-[0-9a-f]{32}\.docx$")
EXAMPLE_CHOICES = {
    "Custom request": None,
    "Demo 1 — Project implementation plan": 0,
    "Demo 2 — Technical and business proposal": 1,
}


def main() -> None:
    st.set_page_config(page_title="Autonomous Document Agent", page_icon="📝", layout="wide")
    st.title("Autonomous Document Agent")
    st.write(
        "Describe the document you need. The agent will create a plan, execute its registered "
        "tools, generate a Word document, and check the result."
    )

    settings = get_settings()
    if not settings.groq_api_key or not settings.groq_model:
        st.warning("Add GROQ_API_KEY and GROQ_MODEL to your local .env file before generating a document.")

    st.selectbox(
        "Start from an example (optional)",
        options=list(EXAMPLE_CHOICES),
        key="example_choice",
        on_change=_load_selected_example,
    )
    with st.form("document_request_form"):
        request_text = st.text_area(
            "What document should I create?",
            key="request_input",
            placeholder="For example: Create a project implementation plan for ...",
            height=180,
            max_chars=settings.max_request_length,
        )
        submitted = st.form_submit_button("Create document", type="primary", use_container_width=True)

    if submitted:
        _run_request(request_text, settings)

    if st.session_state.get("agent_error"):
        st.error(st.session_state["agent_error"])
    result_data = st.session_state.get("agent_result")
    if result_data:
        _show_result(AgentResponse.model_validate(result_data), settings)


def _load_selected_example() -> None:
    example_index = EXAMPLE_CHOICES[st.session_state["example_choice"]]
    if example_index is None:
        return
    examples_path = WORKSPACE / "demo_requests.json"
    examples = json.loads(examples_path.read_text(encoding="utf-8"))
    st.session_state["request_input"] = examples[example_index]["request"]


def _run_request(request_text: str, settings: Settings) -> None:
    st.session_state.pop("agent_result", None)
    st.session_state.pop("agent_error", None)
    try:
        request = request_text.strip()
        if not request:
            raise ValueError("Enter a request before creating a document.")
        if len(request) > settings.max_request_length:
            raise ValueError(f"Your request is longer than the {settings.max_request_length}-character limit.")
        validated = AgentRequest(request=request)
        with st.spinner("Planning, drafting, generating, and checking your document…"):
            result = run_agent(validated.request, settings)
        st.session_state["agent_result"] = result.model_dump(mode="json")
    except ConfigurationError as exc:
        st.session_state["agent_error"] = str(exc)
    except ValueError as exc:
        st.session_state["agent_error"] = str(exc)
    except (PlanningError, RetryableOperationError) as exc:
        logger.warning("Streamlit agent request failed: %s", type(exc).__name__)
        st.session_state["agent_error"] = (
            "The agent could not complete the request. Check the configured Groq model, credentials, "
            "quota, and terminal logs, then try again."
        )
    except Exception:
        logger.exception("Unexpected Streamlit agent failure")
        st.session_state["agent_error"] = "An unexpected error prevented document generation. Check the terminal logs."


def _show_result(result: AgentResponse, settings: Settings) -> None:
    st.subheader("Execution result")
    if result.status == "completed":
        st.success(result.message)
    else:
        st.error(result.message)

    st.caption(f"Execution ID: {result.execution_id} · Document type: {result.document_type}")
    summary = result.execution_summary
    metric_columns = st.columns(4)
    metric_columns[0].metric("Tasks", summary.total_tasks)
    metric_columns[1].metric("Completed", summary.completed_tasks)
    metric_columns[2].metric("Failed", summary.failed_tasks)
    metric_columns[3].metric("Skipped", summary.skipped_tasks)

    st.markdown("#### Agent-generated task list")
    for task in result.plan:
        icon = {
            "completed": "✅",
            "failed": "❌",
            "skipped": "⏭️",
            "running": "⏳",
            "pending": "▫️",
        }.get(task.status.value, "▫️")
        with st.expander(f"{icon} {task.description} — {task.status.value}"):
            st.write(f"**Tool:** `{task.tool_name}`")
            st.write(f"**Expected output:** {task.expected_output}")
            st.write(f"**Dependencies:** {', '.join(task.dependencies) if task.dependencies else 'None'}")
            if task.attempts:
                st.write(f"**Attempts:** {task.attempts}")
            if task.error_message:
                st.error(task.error_message)
            if task.result is not None:
                st.json(task.result)

    st.markdown("#### Document quality checks")
    st.write("Overall:", "✅ Passed" if result.quality_check.passed else "❌ Failed")
    st.dataframe(
        [
            {"Check": check.name.replace("_", " ").title(), "Passed": check.passed, "Details": check.details}
            for check in result.quality_check.checks
        ],
        hide_index=True,
        use_container_width=True,
    )

    if result.document and result.quality_check.passed:
        document_path = _resolve_document_path(result.document.filename, settings)
        if document_path is None:
            st.error("The generated document could not be found in the configured output directory.")
            return
        document_bytes = document_path.read_bytes()
        st.download_button(
            "Download Word document",
            data=document_bytes,
            file_name=result.document.filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
        )
        st.caption(f"{result.document.filename} · {result.document.size_bytes:,} bytes")
        _show_document_preview(document_path)


def _resolve_document_path(filename: str, settings: Settings) -> Path | None:
    if not FILENAME_PATTERN.fullmatch(filename):
        return None
    root = settings.generated_documents_dir.resolve()
    candidate = (root / filename).resolve()
    if candidate.parent != root or not candidate.is_file():
        return None
    return candidate


def _show_document_preview(path: Path) -> None:
    st.markdown("#### Word document preview")
    try:
        document = Document(path)
    except Exception:
        logger.exception("Could not open generated document for Streamlit preview")
        st.warning("The file is available to download, but an in-page preview could not be displayed.")
        return

    with st.container(border=True):
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name == "Title":
                st.title(text)
            elif style_name.startswith("Heading"):
                st.markdown(f"**{text}**")
            elif style_name.startswith("List Bullet"):
                st.markdown(f"- {text}")
            elif style_name.startswith("List Number"):
                st.markdown(f"1. {text}")
            else:
                st.write(text)
        for table in document.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if rows:
                st.table(rows)


if __name__ == "__main__":
    main()
