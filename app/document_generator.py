"""Professional, editable DOCX generation."""

from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas import DocumentDraft

NAVY = "17365D"
BLUE = "2F75B5"


def generate_docx(
    draft: DocumentDraft,
    document_type: str,
    assumptions: list[str],
    output_directory: Path,
) -> Path:
    """Create a uniquely named DOCX inside the configured output directory."""
    output_directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("263238")
    normal.paragraph_format.space_after = Pt(7)
    for style_name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY if style_name != "Heading 2" else BLUE)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(draft.title)
    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(16)
    run = metadata.add_run(f"{document_type}  |  Generated {datetime.now(timezone.utc):%B %d, %Y %H:%M UTC}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("64748B")

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(draft.executive_summary)
    for section_content in draft.sections:
        document.add_heading(section_content.heading, level=1)
        for paragraph in section_content.paragraphs:
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
        for item in section_content.bullets:
            if item.strip():
                document.add_paragraph(item.strip(), style="List Bullet")
        for item in section_content.numbered_items:
            if item.strip():
                document.add_paragraph(item.strip(), style="List Number")
        if section_content.table:
            _add_table(document, section_content.table)

    if assumptions:
        document.add_heading("Assumptions and Constraints", level=1)
        for assumption in assumptions:
            if assumption.strip():
                document.add_paragraph(assumption.strip(), style="List Bullet")
    document.add_heading("Conclusion and Next Steps", level=1)
    document.add_paragraph(draft.conclusion)
    _add_page_number(section)

    output_path = output_directory / f"document-{uuid4().hex}.docx"
    document.save(output_path)
    return output_path


def _add_table(document: Document, rows: list[list[str]]) -> None:
    valid_rows = [row for row in rows if row and any(cell.strip() for cell in row)]
    if not valid_rows:
        return
    column_count = max(len(row) for row in valid_rows)
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(valid_rows):
        cells = table.add_row().cells
        for column_index, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = values[column_index].strip() if column_index < len(values) else ""
            if row_index == 0:
                _shade_cell(cell, BLUE)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)


def _shade_cell(cell: object, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_page_number(section: object) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
