"""Deterministic validation of generated Word documents."""

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.config import Settings, get_settings
from app.schemas import QualityCheck, QualityResult

PLACEHOLDERS = ("todo: fill this in", "[insert text]", "lorem ipsum", "tbd")


def check_document(
    path: str | Path,
    expected_headings: list[str],
    title: str,
    settings: Settings | None = None,
) -> QualityResult:
    """Open the generated artifact and verify content, headings, and size."""
    settings = settings or get_settings()
    file_path = Path(path)
    exists = file_path.is_file()
    size = file_path.stat().st_size if exists else 0
    checks = [
        QualityCheck(name="file_exists", passed=exists, details="The generated file exists." if exists else "The generated file is missing."),
        QualityCheck(
            name="file_size",
            passed=0 < size <= settings.max_document_size_bytes,
            details=f"File size is {size} bytes; maximum allowed is {settings.max_document_size_bytes}.",
        ),
    ]
    extracted_text = ""
    headings: set[str] = set()
    reopened = False
    if exists and size > 0:
        try:
            document = Document(file_path)
            reopened = True
            extracted_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            headings = {
                paragraph.text.strip().casefold()
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name.startswith("Heading")
            }
            extracted_text += "\n" + "\n".join(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
        except (PackageNotFoundError, OSError, ValueError):
            reopened = False
    checks.append(QualityCheck(
        name="docx_reopens",
        passed=reopened,
        details="The file reopened as a DOCX document." if reopened else "The file could not be reopened as a valid DOCX.",
    ))
    text = extracted_text.strip()
    checks.append(QualityCheck(
        name="meaningful_text",
        passed=bool(text) and len(text) >= settings.min_document_characters,
        details=f"Extracted {len(text)} characters; minimum is {settings.min_document_characters}.",
    ))
    checks.append(QualityCheck(
        name="meaningful_title",
        passed=bool(title.strip()) and title.strip().casefold() in text.casefold(),
        details="The document title is present." if title.strip().casefold() in text.casefold() else "The document title is missing.",
    ))
    missing_headings = [heading for heading in expected_headings if heading.strip().casefold() not in headings]
    checks.append(QualityCheck(
        name="expected_headings",
        passed=not missing_headings,
        details="All planned headings are present." if not missing_headings else f"Missing headings: {', '.join(missing_headings)}.",
    ))
    placeholders = [phrase for phrase in PLACEHOLDERS if phrase in text.casefold()]
    checks.append(QualityCheck(
        name="no_placeholders",
        passed=not placeholders,
        details="No obvious placeholders were detected." if not placeholders else f"Placeholder text detected: {', '.join(placeholders)}.",
    ))
    return QualityResult(
        passed=all(check.passed for check in checks),
        checks=checks,
        character_count=len(text),
        file_size_bytes=size,
    )
